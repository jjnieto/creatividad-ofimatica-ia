# Primera versión de matriz_decision.xlsx, la plantilla del Ejercicio 4 (bloque 3).
# Como con las diapositivas, el .xlsx manda: si ya existe, el script no lo sobrescribe.
# Uso: python matriz_decision.py [ruta_salida]
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SALIDA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "matriz_decision.xlsx")

# Código de color del curso: ámbar = trabajo de las personas, verde azulado = IA.
AMBAR, AMBAR_CLARO = "F4A300", "FDEFCC"
VERDE, VERDE_CLARO = "0F7C70", "D5ECE9"
OSCURO, GRIS = "1F2A30", "6B7780"

FUENTE = "Arial"
fino = Side(style="thin", color="BFC5C9")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)

CRITERIOS = ["Impacto", "Viabilidad", "Bajo coste", "Originalidad"]
# La fórmula del guion, con los pesos escritos dentro (paso 5 del Ejercicio 4).
FORMULA = "=B{f}*35%+C{f}*30%+D{f}*20%+E{f}*15%"
EJEMPLO = [
    ("Reserva por internet", 4, 5, 5, 3),
    ("Turno visible en el móvil", 4, 3, 2, 4),
    ("Abono con cita fija", 5, 4, 4, 3),
]


def fuente(**kw):
    return Font(name=FUENTE, size=kw.pop("size", 11), **kw)


def matriz(ws, filas):
    """Tabla en A1:F4 para que la fórmula coincida con la del guion (F2, B2:E2)."""
    for col, texto in enumerate(["Propuesta", *CRITERIOS, "Resultado"], start=1):
        c = ws.cell(row=1, column=col, value=texto)
        c.font = fuente(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=OSCURO)
        c.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center")
        c.border = BORDE
    ws.row_dimensions[1].height = 24
    for f in range(2, 5):
        datos = filas[f - 2] if filas else (None,) * 5
        for col in range(1, 6):
            c = ws.cell(row=f, column=col, value=datos[col - 1])
            c.font = fuente()
            c.fill = PatternFill("solid", fgColor=AMBAR_CLARO)
            c.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center")
            c.border = BORDE
        c = ws.cell(row=f, column=6, value=FORMULA.format(f=f))
        c.font = fuente(bold=True)
        c.number_format = "0.00"
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDE
        ws.row_dimensions[f].height = 22

    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    dv.errorTitle = "Puntuación no válida"
    dv.error = "Puntuad cada criterio con un número entero del 1 al 5."
    dv.promptTitle = "Del 1 al 5"
    dv.prompt = "1 = muy poco, 5 = mucho. En «Bajo coste», 5 = muy económica."
    ws.add_data_validation(dv)
    dv.add("B2:E4")

    ws["D1"].comment = Comment("5 = muy económica; 1 = muy cara.", "Docente")
    ws["F1"].comment = Comment(
        "Se calcula solo: =B2*35%+C2*30%+D2*20%+E2*15%. "
        "Si cambiáis los pesos, hay que cambiarlos en las tres filas.", "Docente")

    ws["A6"] = "Pesos: Impacto 35 % · Viabilidad 30 % · Bajo coste 20 % · Originalidad 15 % (suman 100 %)."
    ws["A6"].font = fuente(size=10, italic=True, color=GRIS)
    ws["A7"] = "La matriz es una ayuda para razonar, no una prueba objetiva: el resultado depende de las puntuaciones y los pesos."
    ws["A7"].font = fuente(size=10, italic=True, color=GRIS)


def anchos(ws):
    ws.column_dimensions["A"].width = 30
    for col in "BCDEF":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["G"].width = 3
    ws.column_dimensions["H"].width = 90


def leyenda(ws, titulo, filas):
    """Una línea corta por fila, sin ajuste de texto: comparte filas con la tabla."""
    ws["H1"] = titulo
    ws["H1"].font = fuente(bold=True, size=12, color=OSCURO)
    for i, texto in enumerate(filas, start=2):
        c = ws.cell(row=i, column=8, value=texto)
        c.font = fuente(size=10)
        c.alignment = Alignment(vertical="center")


def pregunta(ws, fila, texto, color_claro, alto):
    ws.cell(row=fila, column=1, value=texto).font = fuente(bold=True)
    ws.cell(row=fila, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=6)
    for col in range(1, 7):
        ws.cell(row=fila, column=col).border = BORDE
    for col in range(2, 7):
        ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=color_claro)
    ws.cell(row=fila, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=fila, column=2).font = fuente()
    ws.row_dimensions[fila].height = alto


def hoja_matriz(wb):
    ws = wb.active
    ws.title = "Matriz"
    anchos(ws)
    matriz(ws, None)

    ws["A12"] = "Nuestra decisión"
    ws["A12"].font = fuente(bold=True, size=12, color=OSCURO)
    pregunta(ws, 13, "Propuesta ganadora", AMBAR_CLARO, 22)
    pregunta(ws, 14, "Elegimos… porque…", AMBAR_CLARO, 48)
    pregunta(ws, 15, "Si no es la de mayor puntuación, ¿por qué la elegimos?", AMBAR_CLARO, 48)
    pregunta(ws, 16, "Al dar más peso al coste, ¿cambió la clasificación?", AMBAR_CLARO, 48)
    pregunta(ws, 17, "¿Qué observó la IA? ¿Tenía razón?", VERDE_CLARO, 64)

    leyenda(ws, "Cómo usar esta hoja", [
        "1. Escribid vuestras tres propuestas del bloque 2 en la columna «Propuesta».",
        "2. Puntuad cada criterio del 1 al 5 (en «Bajo coste», 5 = muy económica).",
        "    «Resultado» se calcula solo. Tenéis un ejemplo en la hoja «Ejemplo».",
        "3. Ordenad la tabla A1:F4 por «Resultado», de mayor a menor.",
        "4. Dad más peso al coste (los pesos siempre suman 100 %). Están dentro",
        "    de la fórmula: hay que cambiarlos en las tres filas.",
        "5. Pedid a la IA que revise la matriz y comprobad si tiene razón.",
        "6. Decidid y justificad la propuesta ganadora en «Nuestra decisión».",
        "Ámbar claro: lo rellenáis vosotros. Verde azulado: lo que dice la IA.",
        "Sin datos personales ni confidenciales: es una simulación formativa.",
    ])
    ws.freeze_panes = None


def hoja_ejemplo(wb):
    ws = wb.create_sheet("Ejemplo")
    anchos(ws)
    matriz(ws, EJEMPLO)
    leyenda(ws, "Ejemplo", [
        "Propuestas del ejemplo de la peluquería (bloque 3). Los datos son supuestos didácticos.",
        "Comprobación a mano: 4×0,35 + 5×0,30 + 5×0,20 + 3×0,15 = 4,35.",
    ])


def main():
    if os.path.exists(SALIDA):
        sys.exit(f"{SALIDA} ya existe: no se sobrescribe. Pasa otra ruta como argumento.")
    wb = Workbook()
    hoja_matriz(wb)
    hoja_ejemplo(wb)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0
    wb.save(SALIDA)
    print(f"Guardado {SALIDA}")


if __name__ == "__main__":
    main()
